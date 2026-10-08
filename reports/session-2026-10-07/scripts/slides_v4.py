"""Make v4 of the team deck: replace the old UI screenshot on slide 10 with the new UI."""
import shutil, sys
from pptx import Presentation
from pptx.util import Emu
src = "/home/vboxuser/Downloads/Disaster_Response_Intelligence_Repo_Hub_G7_updated_v3.pptx"
out = sys.argv[1]; shot = sys.argv[2]
shutil.copy(src, out)
prs = Presentation(out)
slide = prs.slides[9]
old = next(sh for sh in slide.shapes if sh.shape_type and sh.shape_type.name == "PICTURE")
left, top, width, height = old.left, old.top, old.width, old.height
# keep the frame, fit the new 16:10 screenshot inside it
from PIL import Image
w, h = Image.open(shot).size
scale = min(width / w, height / h)
new_w, new_h = int(w * scale), int(h * scale)
pic = slide.shapes.add_picture(shot, left + (width - new_w) // 2, top + (height - new_h) // 2, new_w, new_h)
# put the new picture where the old one was in z-order, then drop the old one
old._element.addprevious(pic._element)
old._element.getparent().remove(old._element)
prs.save(out)
print("saved", out, "| slide 10 picture replaced:", Emu(new_w).inches, "x", Emu(new_h).inches, "in")
